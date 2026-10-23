import os
from django.shortcuts import get_object_or_404
from django.core.files.storage import default_storage
from django.core.files.base import ContentFile
from rest_framework import status, viewsets, permissions
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser, FormParser

from apps.core.models import Organization
from apps.core.serializers import SuperAdminOrganizationSerializer, OrganizationSettingsSerializer
from apps.accounts.permissions import IsSuperAdmin, IsOrgManager
from apps.loyalty.google_wallet_service import GoogleWalletService
from apps.accounts.permissions import IsSuperAdmin, IsOrgManager

class SuperAdminOrganizationViewSet(viewsets.ModelViewSet):
    """
    ViewSet for global SaaS admin to manage restaurant organizations.
    """
    queryset = Organization.objects.all().order_by('-created_at')
    serializer_class = SuperAdminOrganizationSerializer
    permission_classes = [permissions.IsAuthenticated, IsSuperAdmin]

    def perform_create(self, serializer):
        serializer.save()

    def update(self, request, *args, **kwargs):
        # Allow toggling activation or updating details
        partial = kwargs.pop('partial', False)
        instance = self.get_object_or_404_by_pk(kwargs.get('pk'))
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response(serializer.data)

    def get_object_or_404_by_pk(self, pk):
        return get_object_or_404(Organization, pk=pk)

class OrganizationSettingsView(APIView):
    """
    View for OrgManager to view and modify their organization's settings.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]

    def get(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        serializer = OrganizationSettingsSerializer(org)
        return Response(serializer.data)

    def patch(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        serializer = OrganizationSettingsSerializer(org, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        
        # Trigger webhook registration if bot token is updated
        if 'tg_bot_token' in request.data:
            from apps.loyalty.tasks import register_tg_webhook
            register_tg_webhook.delay(org.id)
            
        return Response(serializer.data)

class OrganizationLogoUploadView(APIView):
    """
    View for OrgManager to upload organization logo image.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        logo_file = request.FILES.get('logo')
        if not logo_file:
            return Response({"error": "No logo file provided"}, status=status.HTTP_400_BAD_REQUEST)

        # Basic extension check
        ext = os.path.splitext(logo_file.name)[1].lower()
        if ext not in ['.jpg', '.jpeg', '.png', '.webp', '.gif']:
            return Response({"error": "Invalid image format"}, status=status.HTTP_400_BAD_REQUEST)

        # Save to media/logos/
        filename = f"logos/logo_{org.slug}{ext}"
        path = default_storage.save(filename, ContentFile(logo_file.read()))
        
        # Build URL
        logo_url = f"/media/{path}"
        
        # Update organization branding config
        branding = org.branding or {}
        branding['logo_url'] = logo_url
        org.branding = branding
        org.save(update_fields=['branding'])

        return Response({
            "status": "success",
            "logo_url": logo_url
        }, status=status.HTTP_200_OK)

class OrganizationBannerUploadView(APIView):
    """
    View for OrgManager to upload organization Google Wallet hero image (banner).
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        banner_file = request.FILES.get('banner') or request.FILES.get('image')
        if not banner_file:
            return Response({"error": "No image file provided"}, status=status.HTTP_400_BAD_REQUEST)

        ext = os.path.splitext(banner_file.name)[1].lower()
        if ext not in ['.jpg', '.jpeg', '.png', '.webp']:
            return Response({"error": "Invalid image format"}, status=status.HTTP_400_BAD_REQUEST)

        filename = f"banners/banner_{org.slug}{ext}"
        path = default_storage.save(filename, ContentFile(banner_file.read()))
        
        banner_url = f"/media/{path}"
        
        # Save to google_wallet_settings
        settings_dict = org.google_wallet_settings or {}
        settings_dict['hero_image_url'] = banner_url
        org.google_wallet_settings = settings_dict
        org.save(update_fields=['google_wallet_settings'])

        return Response({
            "status": "success",
            "banner_url": banner_url
        }, status=status.HTTP_200_OK)

class OrganizationWalletLogoUploadView(APIView):
    """
    View for OrgManager to upload organization Google Wallet logo image.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        logo_file = request.FILES.get('logo') or request.FILES.get('image')
        if not logo_file:
            return Response({"error": "No image file provided"}, status=status.HTTP_400_BAD_REQUEST)

        ext = os.path.splitext(logo_file.name)[1].lower()
        if ext not in ['.jpg', '.jpeg', '.png', '.webp']:
            return Response({"error": "Invalid image format"}, status=status.HTTP_400_BAD_REQUEST)

        filename = f"logos/wallet_logo_{org.slug}{ext}"
        path = default_storage.save(filename, ContentFile(logo_file.read()))
        
        logo_url = f"/media/{path}"
        
        settings_dict = org.google_wallet_settings or {}
        settings_dict['logo_url'] = logo_url
        org.google_wallet_settings = settings_dict
        org.save(update_fields=['google_wallet_settings'])

        return Response({
            "status": "success",
            "logo_url": logo_url
        }, status=status.HTTP_200_OK)

class OrganizationGoogleWalletSettingsView(APIView):
    """
    View for OrgManager to get and setup Google Wallet class.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]

    def get(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        return Response({
            "is_google_wallet_enabled": org.is_google_wallet_enabled,
            "google_issuer_id": org.google_issuer_id,
            "google_loyalty_class_id": org.google_loyalty_class_id,
            "google_wallet_settings": org.google_wallet_settings or {},
        }, status=status.HTTP_200_OK)

    def post(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        
        # Extract and persist all configuration fields first
        issuer_id = request.data.get('issuer_id')
        class_name = request.data.get('class_name')
        issuer_name = request.data.get('issuer_name')
        logo_url = request.data.get('logo_url')
        hero_image_url = request.data.get('hero_image_url')
        hex_background_color = request.data.get('hex_background_color')
        
        settings_dict = org.google_wallet_settings or {}
        if issuer_id is not None:
            settings_dict['issuer_id'] = str(issuer_id).strip()
            org.google_issuer_id = str(issuer_id).strip()
        if class_name is not None:
            settings_dict['class_name'] = str(class_name).strip()
        if issuer_name is not None:
            settings_dict['issuer_name'] = str(issuer_name).strip()
        if logo_url is not None:
            settings_dict['logo_url'] = str(logo_url).strip()
        if hero_image_url is not None:
            settings_dict['hero_image_url'] = str(hero_image_url).strip()
        if hex_background_color is not None:
            settings_dict['hex_background_color'] = str(hex_background_color).strip()
        locations = request.data.get('locations')
        if locations is not None:
            settings_dict['locations'] = locations
        push_templates = request.data.get('push_templates')
        if push_templates is not None:
            settings_dict['push_templates'] = push_templates

        org.google_wallet_settings = settings_dict
        org.save(update_fields=['google_issuer_id', 'google_wallet_settings'])

        if not org.google_issuer_id:
            return Response({"error": "Google Issuer ID is required"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            service = GoogleWalletService()
            result = service.create_or_update_loyalty_class(org, request.data)
            return Response({
                "status": "success",
                "class_id": org.google_loyalty_class_id,
                "google_response": result,
                "google_wallet_settings": org.google_wallet_settings
            })
        except Exception as e:
            return Response({
                "error": str(e),
                "saved_locally": True,
                "google_wallet_settings": org.google_wallet_settings
            }, status=status.HTTP_400_BAD_REQUEST)

class OrganizationWebhookStatusView(APIView):
    """
    View for OrgManager to check Telegram Webhook status.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]

    def get(self, request, organization_id, *args, **kwargs):
        import requests
        from django.conf import settings
        
        org = get_object_or_404(Organization, id=organization_id)
        if not org.tg_bot_token:
            return Response({"error": "Bot token not configured"}, status=status.HTTP_400_BAD_REQUEST)

        expected_url = f"{settings.WEBHOOK_DOMAIN.rstrip('/')}/api/accounts/tma/webhook/{org.slug}/{org.tg_bot_token}/"
        url = f"https://api.telegram.org/bot{org.tg_bot_token}/getWebhookInfo"
        
        try:
            response = requests.get(url, timeout=10)
            response.raise_for_status()
            data = response.json()
            if not data.get('ok'):
                return Response({"error": data.get('description', 'Unknown error')}, status=status.HTTP_400_BAD_REQUEST)
                
            webhook_info = data.get('result', {})
            return Response({
                "expected_url": expected_url,
                "telegram_response": webhook_info
            })
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    def post(self, request, organization_id, *args, **kwargs):
        import requests
        from django.conf import settings
        
        org = get_object_or_404(Organization, id=organization_id)
        if not org.tg_bot_token:
            return Response({"error": "Bot token not configured"}, status=status.HTTP_400_BAD_REQUEST)
            
        webhook_url = f"{settings.WEBHOOK_DOMAIN.rstrip('/')}/api/accounts/tma/webhook/{org.slug}/{org.tg_bot_token}/"
        url = f"https://api.telegram.org/bot{org.tg_bot_token}/setWebhook"
        
        try:
            response = requests.post(url, json={"url": webhook_url}, timeout=10)
            response.raise_for_status()
            data = response.json()
            if not data.get('ok'):
                return Response({"error": data.get('description', 'Unknown error')}, status=status.HTTP_400_BAD_REQUEST)
                
            return Response({"status": "Webhook successfully registered", "expected_url": webhook_url})
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

class OrganizationSendTestMessageView(APIView):
    """
    View for OrgManager to send a test message to verify Bot Token from Organization Settings.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]

    def post(self, request, organization_id, *args, **kwargs):
        import requests
        org = get_object_or_404(Organization, id=organization_id)
        if not org.tg_bot_token:
            return Response({"error": "Бот не настроен для данной организации"}, status=status.HTTP_400_BAD_REQUEST)
            
        telegram_id = request.data.get('telegram_id')
        if not telegram_id:
            return Response({"error": "telegram_id is required"}, status=status.HTTP_400_BAD_REQUEST)
            
        url = f"https://api.telegram.org/bot{org.tg_bot_token}/sendMessage"
        text = f"<b>[ТЕСТ]</b>\nБот заведения <b>{org.name}</b> успешно настроен и подключен!"
        
        try:
            response = requests.post(url, json={
                "chat_id": telegram_id,
                "text": text,
                "parse_mode": "HTML"
            }, timeout=10)
            data = response.json()
            if not data.get('ok'):
                return Response({"error": data.get('description', 'Unknown error')}, status=status.HTTP_400_BAD_REQUEST)
                
                # Auto-save admin's telegram_id to their profile if it changed/was blank
            if request.user.telegram_id != telegram_id:
                request.user.telegram_id = telegram_id
                request.user.save(update_fields=['telegram_id'])
                
            return Response({"status": "success", "detail": "Test message sent successfully"})
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

from apps.core.models import Employee
from apps.core.serializers import TmaEmployeeSerializer

class TmaEmployeeViewSet(viewsets.ModelViewSet):
    """
    ViewSet for OrgManager to manage TMA Employees (Owners, Managers, Cashiers).
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]
    serializer_class = TmaEmployeeSerializer

    def get_queryset(self):
        org_id = self.kwargs.get('organization_id')
        return Employee.objects.filter(organization_id=org_id).order_by('-id')

    def perform_create(self, serializer):
        org_id = self.kwargs.get('organization_id')
        org = get_object_or_404(Organization, id=org_id)
        serializer.save(organization=org)

class OrganizationTestIikoConnectionView(APIView):
    """
    View for OrgManager to test iiko credentials and verify token acquisition.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]

    def post(self, request, organization_id, *args, **kwargs):
        from apps.loyalty.services import IikoAuthService
        from django.core.cache import cache
        
        org = get_object_or_404(Organization, id=organization_id)
        
        # Override with request data if transient values provided
        iiko_api_login = request.data.get('iiko_api_login', org.iiko_api_login)
        iiko_app_id = request.data.get('iiko_app_id', org.iiko_app_id)
        iiko_client_secret = request.data.get('iiko_client_secret', org.iiko_client_secret)
        iiko_api_base_url = request.data.get('iiko_api_base_url', org.iiko_api_base_url)
        
        if not iiko_api_login and not iiko_app_id:
            return Response({"error": "Укажите iiko API Login (apiKey) или App ID"}, status=status.HTTP_400_BAD_REQUEST)

        test_org = Organization(
            id=org.id,
            iiko_api_base_url=iiko_api_base_url or "https://api-ru.iiko.services/api/1",
            iiko_api_login=iiko_api_login or "",
            iiko_app_id=iiko_app_id or "",
            iiko_client_secret=iiko_client_secret or ""
        )
        
        try:
            cache.delete(f"iiko_token_{org.id}")
        except Exception:
            pass
        
        try:
            auth_service = IikoAuthService(test_org)
            token = auth_service.get_access_token()
            return Response({
                "status": "success",
                "message": "Подключение к iiko успешно установлено! Токен авторизации получен.",
                "token_prefix": f"{token[:8]}..." if len(token) > 8 else token
            })
        except Exception as e:
            return Response({"error": f"Ошибка авторизации iiko: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)


class OrganizationTestIikoServerConnectionView(APIView):
    """
    View for OrgManager to test connection to iiko Server (RMS / Chain) and verify auth.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]

    def post(self, request, organization_id, *args, **kwargs):
        from apps.loyalty.services import IikoServerAuthService
        from django.core.cache import cache
        
        org = get_object_or_404(Organization, id=organization_id)
        
        # Override with request data if transient values provided
        iiko_server_url = request.data.get('iiko_server_url', org.iiko_server_url)
        iiko_server_login = request.data.get('iiko_server_login', org.iiko_server_login)
        iiko_server_password = request.data.get('iiko_server_password', org.iiko_server_password)
        
        # If the client sent the masked password placeholder, fetch the saved password from db
        if iiko_server_password == '••••••••':
            iiko_server_password = org.iiko_server_password

        if not iiko_server_url:
            return Response({"error": "Укажите адрес сервера (iiko Server URL)"}, status=status.HTTP_400_BAD_REQUEST)
        if not iiko_server_login or not iiko_server_password:
            return Response({"error": "Укажите логин и пароль пользователя сервера"}, status=status.HTTP_400_BAD_REQUEST)

        test_org = Organization(
            id=org.id,
            iiko_server_url=iiko_server_url,
            iiko_server_login=iiko_server_login,
            iiko_server_password=iiko_server_password
        )
        
        # Invalidate existing server token cache to force a fresh test authentication
        try:
            cache_key = f"iiko_server_token_{org.id}"
            cache.delete(cache_key)
        except Exception:
            pass
        
        try:
            auth_service = IikoServerAuthService(test_org)
            token = auth_service.get_access_token()
            
            # Since testing connections consumes a slot, we should immediately release it
            try:
                auth_service.logout()
            except Exception:
                pass
                
            return Response({
                "status": "success",
                "message": "Подключение к iiko Server успешно установлено! Авторизационный токен получен.",
                "token_prefix": f"{token[:8]}..." if len(token) > 8 else token
            })
        except Exception as e:
            return Response({"error": f"Ошибка авторизации на сервере iiko: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)


from apps.core.models import WabaTriggerRule, WabaMessageLog
from apps.core.serializers import WabaTriggerRuleSerializer, WabaMessageLogSerializer
from apps.core.services.waba_service import WabaService

class OrganizationWabaTemplatesView(APIView):
    """
    Fetch approved WABA message templates directly from Meta Graph API.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]

    def get(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        if not org.waba_access_token or not org.waba_business_account_id:
            return Response(
                {"error": "WABA Access Token и Business Account ID должны быть заполнены и сохранены."},
                status=status.HTTP_400_BAD_REQUEST
            )
        service = WabaService(org)
        try:
            templates = service.fetch_templates()
            return Response({"templates": templates})
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)


class OrganizationWabaPhoneNumbersView(APIView):
    """
    Fetch registered phone numbers associated with the WABA account.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]

    def get(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        if not org.waba_access_token or not org.waba_business_account_id:
            return Response(
                {"error": "WABA Access Token и Business Account ID должны быть заполнены и сохранены."},
                status=status.HTTP_400_BAD_REQUEST
            )
        service = WabaService(org)
        try:
            numbers = service.fetch_phone_numbers()
            return Response({"phone_numbers": numbers})
        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)


class OrganizationWabaRulesViewSet(viewsets.ModelViewSet):
    """
    Manage trigger rules for automatic WABA notifications.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]
    serializer_class = WabaTriggerRuleSerializer

    def get_queryset(self):
        org_id = self.kwargs.get('organization_id')
        return WabaTriggerRule.objects.filter(organization_id=org_id).order_by('id')

    def perform_create(self, serializer):
        org_id = self.kwargs.get('organization_id')
        org = get_object_or_404(Organization, id=org_id)
        serializer.save(organization=org)


class OrganizationWabaSendTestView(APIView):
    """
    Send a test message using a configured rule or template and mappings.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]

    def post(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        if not org.waba_access_token or not org.waba_phone_number_id:
            return Response(
                {"error": "Укажите и сохраните WABA Access Token и Phone Number ID перед тестированием."},
                status=status.HTTP_400_BAD_REQUEST
            )

        phone = request.data.get('phone')
        if not phone:
            return Response({"error": "Укажите номер телефона для тестовой отправки"}, status=status.HTTP_400_BAD_REQUEST)

        template_name = request.data.get('template_name')
        if not template_name:
            return Response({"error": "Укажите имя шаблона"}, status=status.HTTP_400_BAD_REQUEST)

        language_code = request.data.get('template_language', 'ru')
        parameter_mappings = request.data.get('parameter_mappings', {})
        sample_data = request.data.get('sample_data', {})

        mock_payload = {
            "phone": phone,
            "text": sample_data.get('text', "Ваш проверочный код: 4921"),
            "sum": sample_data.get('sum', 500),
            "balance": sample_data.get('balance', 1500),
        }

        service = WabaService(org)
        components, rendered_text = service.build_components(parameter_mappings, mock_payload)

        res = service.send_template_message(
            to_phone=phone,
            template_name=template_name,
            language_code=language_code,
            components=components,
            message_type='test',
            rendered_text=rendered_text
        )

        if res.get('success'):
            return Response({
                "status": "success",
                "message": "Тестовое сообщение успешно отправлено!",
                "message_id": res.get('message_id'),
                "rendered_text": rendered_text
            })
        else:
            return Response({
                "error": res.get('error', 'Неизвестная ошибка при отправке'),
                "rendered_text": rendered_text
            }, status=status.HTTP_400_BAD_REQUEST)


class OrganizationWabaLogsView(APIView):
    """
    View dispatch logs for WABA messages of an organization with filtering and Excel export.
    """
    permission_classes = [permissions.IsAuthenticated, IsOrgManager]

    def get(self, request, organization_id, *args, **kwargs):
        org = get_object_or_404(Organization, id=organization_id)
        qs = WabaMessageLog.objects.filter(organization=org)

        # 1. Date filters
        start_date_str = request.query_params.get('start_date')
        if start_date_str:
            try:
                from datetime import datetime
                sd = datetime.strptime(start_date_str[:10], '%Y-%m-%d')
                qs = qs.filter(created_at__gte=sd)
            except ValueError:
                pass

        end_date_str = request.query_params.get('end_date')
        if end_date_str:
            try:
                from datetime import datetime, time
                ed = datetime.strptime(end_date_str[:10], '%Y-%m-%d')
                ed_end = datetime.combine(ed.date(), time.max)
                qs = qs.filter(created_at__lte=ed_end)
            except ValueError:
                pass

        # 2. Status filter
        status_filter = request.query_params.get('status')
        if status_filter and status_filter in ['SENT', 'FAILED', 'PENDING']:
            qs = qs.filter(status=status_filter)

        # 3. Search query
        search_query = request.query_params.get('search', '').strip()
        if search_query:
            from django.db.models import Q
            qs = qs.filter(
                Q(phone__icontains=search_query) |
                Q(template_name__icontains=search_query) |
                Q(rendered_text__icontains=search_query) |
                Q(waba_message_id__icontains=search_query) |
                Q(error_message__icontains=search_query)
            )

        qs = qs.order_by('-created_at')

        # 4. Check if export is requested
        export_format = request.query_params.get('export')
        if export_format in ['excel', 'xlsx']:
            import io
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
            from django.http import HttpResponse

            export_qs = qs[:10000]

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "WABA Logs"

            headers = [
                "ID", "Дата и время", "Телефон", "Событие", 
                "Шаблон", "Язык", "Статус", "WABA Message ID", 
                "Текст сообщения", "Текст ошибки"
            ]
            ws.append(headers)

            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            event_labels = {
                'auth_code': 'Код авторизации',
                'balance': 'Баланс кошелька',
                'push': 'Информационный пуш',
                'custom': 'Пользовательский'
            }

            status_labels = {
                'SENT': 'Отправлено',
                'FAILED': 'Ошибка',
                'PENDING': 'В очереди'
            }

            for row_idx, item in enumerate(export_qs, start=2):
                created_str = item.created_at.strftime('%Y-%m-%d %H:%M:%S') if item.created_at else ''
                status_display = status_labels.get(item.status, item.status)
                event_display = event_labels.get(item.message_type, item.message_type)

                ws.append([
                    item.id,
                    created_str,
                    item.phone or '',
                    event_display,
                    item.template_name or '',
                    item.language or '',
                    status_display,
                    item.waba_message_id or '',
                    item.rendered_text or '',
                    item.error_message or ''
                ])

                status_cell = ws.cell(row=row_idx, column=7)
                if item.status == 'SENT':
                    status_cell.font = Font(color="166534", bold=True)
                elif item.status == 'FAILED':
                    status_cell.font = Font(color="991B1B", bold=True)

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 60)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            start_label = start_date_str or 'all'
            end_label = end_date_str or 'now'
            filename = f"waba_logs_{org.slug}_{start_label}_to_{end_label}.xlsx"
            response = HttpResponse(
                output.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response

        limit = min(int(request.query_params.get('limit', 100)), 500)
        logs = qs[:limit]
        serializer = WabaMessageLogSerializer(logs, many=True)
        return Response({
            "logs": serializer.data,
            "total_count": qs.count()
        })



